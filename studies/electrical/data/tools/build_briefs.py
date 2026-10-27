"""Briefs for the 17 initial ideas (2026-10-27).

Drafts by three helpers (ideation/briefs-work/draft-h*.json: reference transcript
via NexLev + web-checked facts), each re-checked by a second helper
(check-c*.json: corrected_brief). This joins the corrected briefs to
initial-ideas.csv and writes ideation/briefs-2026-10-27.tsv and .xlsx
(Title / Reference / Brief, plus Idea and Sources).
"""
import csv, json, glob, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
IDEA = os.path.join(HERE, '..', '..', 'ideation')
W = os.path.join(IDEA, 'briefs-work')

drafts = {b['key']: b for f in glob.glob(os.path.join(W, 'draft-h*.json')) for b in json.load(open(f))}
checks = {b['key']: b for f in glob.glob(os.path.join(W, 'check-c*.json')) for b in json.load(open(f))}
rows = []
for x in csv.DictReader(open(os.path.join(IDEA, 'initial-ideas.csv'), encoding='utf-8')):
    k = x['key']; brief = checks[k]['corrected_brief'].strip()
    assert '\n' not in brief and 200 <= len(brief.split()) <= 300 and '$' not in brief, k
    ref = f"\"{x['ref_title']}\" · {int(x['ref_views']):,} ({x['ref_channel']}, {x['ref_date']}) · https://www.youtube.com/watch?v={x['ref_id']}"
    srcs = ' | '.join(drafts[k].get('sources', []))
    rows.append([k, x['title'], ref, brief, srcs])
hdr = ['Idea', 'Title', 'Reference', 'Brief', 'Sources (facts checked against)']
with open(os.path.join(IDEA, 'briefs-2026-10-27.tsv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f, delimiter='\t'); w.writerow(hdr[1:4]); w.writerows([r[1:4] for r in rows])
wb = Workbook(); ws = wb.active; ws.title = 'Briefs'
ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True)
for r in rows: ws.append(r)
for col, wd in zip('ABCDE', [6, 40, 45, 110, 60]): ws.column_dimensions[col].width = wd
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(os.path.join(IDEA, 'briefs-2026-10-27.xlsx'))
print(len(rows), 'briefs')
