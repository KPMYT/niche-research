"""Thumbnail instructions for the 17 initial ideas (operator, 2026-10-27).

Reads ideation/initial-ideas.csv for titles and references; the instructions are
the operator-approved thumbnail advice, rewritten short. A1 follows the
operator's own call: keep the reference with its 2 arrows, change only the
text and the character, remove the logo.
Writes ideation/thumb-instructions.csv and .xlsx.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
IDEA = os.path.join(HERE, '..', '..', 'ideation')

# key: (instructions, style reference id or '', what the style reference is for)
INSTR = {
 'A1': ("Recreate the reference as it is, with its 2 curved arrows (From Here -> To Here). "
        "Only changes: our character in place of the man, kneeling at the wall marking the new box; "
        "the top banner reads NO PATCHING; remove the green logo bottom left.", '', ''),
 'A2': ("Keep the reference close-up: short wires sticking out of the blue box with the tape measure under them. "
        "Add our character's face at the side, frowning at the box. Text TOO SHORT? in place of 'Too Short'. Remove the logo.", '', ''),
 'A7': ("Blue explainer board: a blueprint drawing of panel -> interlock kit -> inlet box -> yellow cord -> generator, one line joining them. "
        "Our character beside the board, pointing at the interlock. Text THE LEGAL WAY.", '', ''),
 'A12': ("Close-up of a plug sagging out of a worn outlet, a light brown scorch mark on the face plate (no flames). "
         "Our character pointing at the gap. One line of text only: REPLACE THIS. Remove the logo.", '', ''),
 'A14': ("A wall switch clearly in the OFF position in front, a bulb faintly glowing behind it in a dim room. "
         "Our character puzzled, pointing at the bulb. Text SWITCH IS OFF?!", '', ''),
 'A29': ("Do not copy the reference layout. An outlet pulled out of the box, a backstabbed wire with a small brown scorch (no flames), "
         "one red circle around it, our character pointing at it. Text CHECK YOURS.", '', ''),
 'A30': ("Real-looking close-up of an open panel, one double-tapped breaker (two wires under one screw) circled in red, "
         "our character's finger on it. Text CHECK YOUR PANEL.", 'nIfmEsPi48c', 'finger on the breaker, real panel photo (Backyard Maine, 805K)'),
 'A31': ("A yellow double-male cord plugged into a dryer outlet, our character holding up a hand to say stop, "
         "a small lineman up a pole in the background. Text STOP DOING THIS.", '', ''),
 'A33': ("Blue explainer board: a blueprint of a house with numbered spots (panel, GFCI, outlets, switches), our character beside it. "
         "Text KNOW THESE FIRST. Backup option: a backstabbed outlet under a red 'no' circle with our character looking doubtful, as in the reference.", '', ''),
 'B3': ("Keep the reference's close-up of a plug with a padlock through the prong holes; add our character at the side pointing at the holes. "
        "Text WHY THE HOLES? (3 words, not his 9). Backup option: a board cutaway of the outlet contacts clicking into the holes.", '', ''),
 'B6': ("One outlet with three numbered red circles (1 backstab, 2 loose loop, 3 nicked wire), our character pointing at number 1. "
        "Text 3 MISTAKES.", '', ''),
 'B7': ("Close-up of a clear solder-seal connector shrinking onto a wire under a heat gun, a soldering iron crossed out at the side, "
        "our character's hand holding the heat gun. Text NO SOLDER NEEDED.", '', ''),
 'B10': ("Blue blueprint background, a cutaway of metal conduit with Romex inside, a big red question mark, "
         "our character with a hand on his chin. Text ILLEGAL?", '', ''),
 'B11': ("The yellow double-male cord hanging on a store hook with a price tag, our character pointing at it. "
         "Text KILLER CORD. (Different from A31, where the cord is plugged in.)", '', ''),
 'C2': ("Split picture: top = wrong (a wire stretched into a wire nut, red X); bottom = right (a short pigtail with a lever connector, green check). "
        "Our character small in a corner. No text.", 'NDqheegj6k8', 'red X / green check split (Handcraft, 10.8M)'),
 'C12': ("Our character holding a worn outlet just pulled out of the wall, a loose plug on the floor, a red arrow to the outlet. "
         "Text PULL THIS OUT.", 'rNeCxWUXOgA', "character holding the thing pulled out, red arrow (Bertha's Clean Home, 710K)"),
 'C15': ("Two spots joined by a dashed line: a dead outlet in the living room and a tripped GFCI in the bathroom, "
         "our character pointing at the GFCI. Text NOT THE BREAKER.", '', ''),
}

def vid(i): return f'https://www.youtube.com/watch?v={i}' if i else ''
def img(i): return f'https://i.ytimg.com/vi/{i}/maxresdefault.jpg' if i else ''

rows = []
for x in csv.DictReader(open(os.path.join(IDEA, 'initial-ideas.csv'), encoding='utf-8')):
    k = x['key']; ins, sid, sfor = INSTR[k]
    rows.append({
        'Idea': k, 'Title': x['title'], 'Look': {'Point': 'Pointing', 'Board': 'Explainer board'}[x['thumb_look']],
        'Thumbnail text': x['thumb_text'] or '(no text)', 'Instructions': ins,
        'Reference video': vid(x['ref_id']), 'Reference thumbnail': img(x['ref_id']),
        'Reference': f"{x['ref_title']} | {x['ref_channel']} | {int(x['ref_views']):,} views",
        'Style reference video': vid(sid), 'Style reference thumbnail': img(sid), 'Style reference is for': sfor,
    })
cols = list(rows[0])
with open(os.path.join(IDEA, 'thumb-instructions.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, cols); w.writeheader(); w.writerows(rows)

wb = Workbook(); ws = wb.active; ws.title = 'Thumbnails'
ws.append(['All thumbnails: our character (channel still to be decided); remove any channel logos from references; no flames.'])
ws.append(cols)
for c in ws[2]: c.font = Font(bold=True)
for r in rows: ws.append([r[c] for c in cols])
widths = {'Idea': 6, 'Title': 45, 'Look': 14, 'Thumbnail text': 18, 'Instructions': 70, 'Reference video': 30,
          'Reference thumbnail': 30, 'Reference': 45, 'Style reference video': 30, 'Style reference thumbnail': 30, 'Style reference is for': 35}
for i, c in enumerate(cols):
    ws.column_dimensions[chr(65 + i)].width = widths[c]
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if isinstance(c.value, str) and c.value.startswith('https://'):
            c.hyperlink = c.value; c.font = Font(color='0563C1', underline='single')
wb.save(os.path.join(IDEA, 'thumb-instructions.xlsx'))
print(len(rows), 'rows')
